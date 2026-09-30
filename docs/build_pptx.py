"""Build docs/Presentasi_Tugas1.pptx from content.py  (python docs/build_pptx.py)"""
import pathlib, sys, math
import matplotlib
from PIL import Image as PILImage, ImageFont

D = pathlib.Path(__file__).parent
sys.path.insert(0, str(D))
import content as K
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY, TEAL, ORANGE, GREY = "#000000", "#000000", "#333333", "#333333"
FD = pathlib.Path(matplotlib.get_data_path()) / "fonts/ttf"
_fonts = {}
def font(px, bold=False):
    k = (round(px * 4), bold)
    if k not in _fonts:
        _fonts[k] = ImageFont.truetype(str(FD / ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")), max(1, int(round(px * 4))))
    return _fonts[k]

def text_w(t, pt, bold=False): return font(pt, bold).getlength(t) / 4.0     # in pt (at 72 dpi 1 px = 1 pt)

def wrap_lines(t, pt, width_pt, bold=False):
    lines = 0
    for para in str(t).split("\\n"):
        words, cur, n = para.split(" "), "", 1
        for w in words:
            trial = (cur + " " + w).strip()
            if text_w(trial, pt, bold) <= width_pt or not cur:
                cur = trial
            else:
                n += 1; cur = w
        lines += n
    return lines

rgb = lambda h: RGBColor.from_string(h.lstrip("#"))
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SN = [0]

def new_slide(title):
    sl = prs.slides.add_slide(BLANK); SN[0] += 1
    ln = sl.shapes.add_shape(1, Inches(.45), Inches(0.92), Inches(12.4), Inches(0.03)); ln.fill.solid(); ln.fill.fore_color.rgb = rgb("#000000"); ln.line.fill.background()
    add_text(sl, .45, .12, 12.4, .75, [title], 25 if len(title) < 60 else 21, True, "#000000", min_pt=16, anchor="mid")
    add_text(sl, .45, 7.08, 11, .3, [f"{K.TITLE}  |  {K.NAMA}"], 9, False, "#333333", min_pt=8)
    add_text(sl, 12.2, 7.08, .8, .3, [str(SN[0])], 9, False, "#333333", min_pt=8, align=PP_ALIGN.RIGHT)
    return sl

def add_text(sl, x, y, w, h, paras, max_pt=16, bold=False, color="#000000", min_pt=9, space=4, bullets=False, align=None, anchor=None):
    """paras: list of str or (str, dict(bold=..., color=...)). Picks the largest font that fits."""
    items = [(p, {}) if isinstance(p, str) else p for p in paras]
    inner_w = (w - 0.2) * 72; inner_h = (h - 0.1) * 72
    pt = max_pt
    while pt > min_pt:
        tot = 0
        for t, o in items:
            t2 = ("• " + t) if bullets else t
            tot += wrap_lines(t2, pt, inner_w - (12 if bullets else 0), o.get("bold", bold)) * pt * 1.2 + space
        if tot <= inner_h:
            break
        pt -= 0.5
    box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05)
    if anchor == "mid":
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (t, o) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " + t) if bullets else t
        p.font.size = Pt(pt); p.font.bold = o.get("bold", bold); p.font.color.rgb = rgb(o.get("color", color)); p.font.name = "Calibri"
        p.space_after = Pt(space)
        if align: p.alignment = align
    return pt

def add_table(sl, rows, x, y, widths, max_h, max_pt=12, min_pt=6.5, first_bold=False):
    tw = sum(widths)
    pt = max_pt
    def heights(pt):
        hs = []
        for i, r in enumerate(rows):
            lines = max(wrap_lines(str(c), pt, (widths[j] - 0.16) * 72, bold=(i == 0)) for j, c in enumerate(r))
            hs.append((lines * pt * 1.2 + 7) / 72.0)
        return hs
    while pt > min_pt and sum(heights(pt)) > max_h:
        pt -= 0.5
    hs = heights(pt)
    shp = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(tw), Inches(sum(hs)))
    t = shp.table
    for j, w_ in enumerate(widths): t.columns[j].width = Inches(w_)
    for i, r in enumerate(rows):
        t.rows[i].height = Inches(hs[i])
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(pt); p.font.name = "Calibri"; p.font.bold = (i == 0) or (first_bold and j == 0)
                p.font.color.rgb = rgb("#000000")
            c.fill.solid(); c.fill.fore_color.rgb = rgb("#E5E7EB" if i == 0 else "#FFFFFF")
    return pt, sum(hs)

def add_pic(sl, path, x, y, w, h):
    iw, ih = PILImage.open(path).size
    sc = min(w / iw, h / ih); pw, ph = iw * sc, ih * sc
    sl.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph))

FIG = lambda n: K.FIGDIR / n
TMP = D / "_formulas"

# ------------------------------------------------------------------ slides
s0 = prs.slides.add_slide(BLANK); SN[0] += 1
bg = s0.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = rgb("#FFFFFF"); bg.line.fill.background()
add_text(s0, .8, 1.5, 11.7, 1.8, [K.TITLE], 38, True, "#000000", min_pt=26)
add_text(s0, .8, 3.4, 11.7, 1.0, [K.SUBTITLE], 20, False, "#333333", min_pt=14)
add_text(s0, .8, 4.6, 11.7, 1.8, ["Laporan Progres Tugas 1", K.COURSE, f"{K.NAMA}  |  NIM {K.NIM}  |  pengerjaan individu", K.REPO], 15, False, "#000000", min_pt=11)

sl = new_slide("Ringkasan eksekutif")
kp = [["Metrik (300 skenario, per bulan)"] + [K.ARM_NAME[a] for a in K.ARMS_K]]
kp.append(["Margin bersih (USD)"] + [K.usd(K.mv(a, "margin")) for a in K.ARMS_K])
kp.append(["Tepat waktu (OTIF)"] + [K.pc(K.mv(a, "otif"), 0) for a in K.ARMS_K])
kp.append(["Waktu ke penawaran pertama (jam)"] + [K.n(K.mv(a, "ttq_h"), 1) for a in K.ARMS_K])
kp.append(["Sentuhan manusia per order"] + [K.n(K.mv(a, "touches_per_order"), 1) for a in K.ARMS_K])
kp.append(["Dana dialihkan serangan (USD)"] + [K.n(K.mv(a, "diverted"), 0, "", True) for a in K.ARMS_K])
add_table(sl, kp, .45, 1.1, [3.8, 2.0, 2.2, 2.6, 2.4], 2.4, 12.5)
add_text(sl, .45, 3.7, 12.4, 3.3, K.EXEC[3:], 12.5, False, "#000000", min_pt=8, bullets=True, space=6)
sl = new_slide("Peta jawaban Tugas 1")
add_table(sl, [["Instruksi", "Jawaban", "Lokasi"]] + [list(r) for r in K.TUGAS_MAP], .45, 1.15, [3.6, 7.4, 1.4], 5.8, 12)
sl = new_slide("1. Identitas")
add_table(sl, K.IDENT, .45, 1.2, [3.2, 9.2], 5.5, 13)
sl = new_slide("2. Deskripsi problem")
add_text(sl, .45, 1.1, 6.1, 5.9, K.PROBLEM[:2], 13, False, "#000000", min_pt=8, bullets=True)
add_text(sl, 6.75, 1.1, 6.1, 5.9, K.PROBLEM[2:], 13, False, "#000000", min_pt=8, bullets=True)
sl = new_slide("2. Lingkungan dan PEAS")
add_table(sl, K.ENV_TABLE, .45, 1.1, [2.0, 2.0, 8.4], 2.4, 11)
add_table(sl, K.PEAS, .45, 3.7, [2.0, 10.4], 3.3, 11)
sl = new_slide("3.1 SOTA: tiga penelitian serupa")
add_table(sl, [["Penelitian", "Temuan", "Relevansi dan celah", "DOI"]] + [[a.split(" doi:")[0], b, c, a.split("doi:")[1]] for a, b, c, _ in K.RELATED_PROBLEM], .35, 1.1, [4.2, 3.0, 3.9, 1.6], 5.7, 10.5)
sl = new_slide("3.2 SOTA: tiga penelitian yang memakai agen cerdas")
add_table(sl, [["Penelitian", "Temuan", "Relevansi dan celah", "DOI"]] + [[a.split(" doi:")[0], b, c, a.split("doi:")[1]] for a, b, c, _ in K.RELATED_AGENT], .35, 1.1, [4.2, 3.0, 3.9, 1.6], 5.7, 10.5)
sl = new_slide("3.3 SOTA lengkap (1/2): agen, negosiasi, pembelajaran")
add_table(sl, [["Referensi", "Jurnal (kuartil*)", "DOI"]] + [[r[2], f"{r[3]} ({r[4]})", r[5]] for r in K.SOTA_ROWS if r[1] == "B"], .35, 1.1, [7.4, 2.6, 2.6], 5.4, 10)
add_text(sl, .35, 6.55, 12.6, .5, [K.CITE_NOTE], 8, False, GREY, min_pt=6.5)
sl = new_slide("3.3 SOTA lengkap (2/2): logistik ekspor, penjadwalan, UMKM digital")
add_table(sl, [["Referensi", "Jurnal (kuartil*)", "DOI"]] + [[r[2], f"{r[3]} ({r[4]})", r[5]] for r in K.SOTA_ROWS if r[1] == "A"], .35, 1.1, [7.4, 2.6, 2.6], 5.4, 10)
add_text(sl, .35, 6.55, 12.6, .5, ["*Kuartil: perkiraan SJR, perlu dicek di scimagojr.com. Semua DOI diverifikasi melalui pencarian web."], 8.5, False, GREY, min_pt=6.5)
sl = new_slide("3.4 Celah penelitian dan kontribusi KCMAS")
add_text(sl, .45, 1.1, 12.4, 5.9, K.SOTA_GAP, 14, False, "#000000", min_pt=9, bullets=True, space=8)
sl = new_slide("4. Tujuan proyek")
add_text(sl, .45, 1.05, 12.4, .85, [K.TUJUAN_UMUM], 12.5, False, "#000000", min_pt=9)
add_table(sl, K.TUJUAN, .35, 1.95, [.8, 4.4, 3.4, 4.0], 3.3, 10.5)
add_text(sl, .45, 5.6, 12.4, 1.4, K.HIPOTESIS, 11.5, False, "#000000", min_pt=8, bullets=True)
sl = new_slide("5.1 Alur end-to-end")
add_pic(sl, FIG("fig_flow.png"), .3, .98, 12.7, 5.85)
add_text(sl, .45, 6.75, 12.4, .35, ["Kotak abu-abu = keputusan pemilik; garis putus-putus = pengecualian; garis titik = jalur paralel."], 10.5, False, GREY, min_pt=8)
sl = new_slide("5.1 Alur end-to-end: langkah 1-6")
add_text(sl, .45, 1.0, 12.4, .8, [K.FLOW_INTRO], 11, False, "#374151", min_pt=8)
add_table(sl, [K.FLOW_STEPS[0]] + K.FLOW_STEPS[1:7], .3, 1.85, [.5, 1.9, 5.2, 3.2, 1.7], 5.15, 10)
sl = new_slide("5.1 Alur end-to-end: langkah 7-12 dan cara membacanya")
add_table(sl, [K.FLOW_STEPS[0]] + K.FLOW_STEPS[7:], .3, 1.05, [.5, 1.9, 5.2, 3.2, 1.7], 3.6, 10)
add_text(sl, .45, 4.85, 12.4, 2.2, K.FLOW_WHY, 11, False, "#000000", min_pt=8, bullets=True, space=5)
sl = new_slide("5.1 Arsitektur KCMAS")
add_pic(sl, FIG("fig_architecture.png"), .4, 1.0, 12.5, 5.3)
add_text(sl, .45, 6.3, 12.4, .75, [K.EVENTBUS_TEXT], 10.5, False, "#000000", min_pt=8)
sl = new_slide("5.1 Siklus agen, tangga konsesi, dan pemasok gagal")
add_pic(sl, FIG("fig_cycle.png"), .3, 1.05, 12.7, 1.6); add_pic(sl, FIG("fig_ladder.png"), .3, 2.7, 6.2, 4.3); add_pic(sl, FIG("fig_failure.png"), 6.6, 2.9, 6.4, 4.0)
sl = new_slide("5.2 Agen dan komponen internal (1/2)")
add_table(sl, K.AGENTS[:6], .25, 1.05, [1.4, 1.2, 2.9, 2.9, 1.8, 1.4, 1.3], 5.9, 10)
sl = new_slide("5.2 Agen dan komponen internal (2/2)")
add_table(sl, [K.AGENTS[0]] + K.AGENTS[6:], .25, 1.05, [1.4, 1.2, 2.9, 2.9, 1.8, 1.4, 1.3], 5.9, 10.5)
sl = new_slide("5.2 Peta agen ke kode")
add_table(sl, K.CODE_MAP, .3, 1.05, [2.2, 3.0, 3.6, 3.8], 5.9, 10)
sl = new_slide("5.2 AI, ML, DL, atau LLM? Alasan pemilihan")
add_table(sl, K.PORTFOLIO, .25, 1.05, [1.8, 2.3, 1.3, 3.9, 3.7], 5.4, 10)
add_text(sl, .45, 6.5, 12.4, .55, [K.CRIT_NOTE], 9.5, False, GREY, min_pt=7)
sl = new_slide("5.3 Kontrak input dan output pengguna")
add_table(sl, K.KONTRAK_IN, .3, 1.05, [1.4, 4.0, 4.6, 2.6], 3.6, 10)
add_table(sl, K.KONTRAK_OUT, .3, 4.8, [1.8, 7.6, 3.2], 2.2, 10)
sl = new_slide("5.4 Menghindari cognitive overload")
add_table(sl, [["Prinsip", "Penerapan pada KCMAS"]] + [list(x) for x in K.COGNITIVE], .35, 1.1, [3.0, 9.6], 5.8, 11.5)
sl = new_slide("5.5 Jadwal, approval, dan eskalasi")
add_table(sl, K.JADWAL, .3, 1.05, [3.0, 4.0, 3.8, 1.8], 3.3, 10)
add_table(sl, K.APPROVAL_LEVELS, .3, 4.6, [2.4, 6.0, 4.2], 2.4, 10)
sl = new_slide("5.6 Koordinasi, negosiasi, dan keamanan")
add_table(sl, [["Mekanisme", "Cara kerja pada KCMAS"]] + [list(x) for x in K.NEGO], .35, 1.1, [3.4, 9.2], 5.8, 11.5)
sl = new_slide("5.6 Uji serangan pada bus pesan")
add_table(sl, K.ATTACKS, .35, 1.1, [4.6, 8.0], 3.0, 11)
add_text(sl, .45, 4.3, 12.4, 2.6, ["Scout agent mobile: desain dan demo aturan migrasi (trust >= 0,80, risiko <= 0,20, state terverifikasi); tidak dipakai oleh simulator dan aplikasi.", "Tes kode aplikasi: lantai harga dan negosiasi tidak berubah oleh teks RFQ; draf tidak membawa harga pemasok."], 12, False, "#000000", min_pt=8, bullets=True)
sl = new_slide("5.7 Mengapa multi-agent, bukan single agent?")
add_table(sl, K.CRITERIA, .45, 1.1, [2.6, 3.8, 3.8, 2.2], 2.3, 10.5)
add_text(sl, .45, 3.55, 12.4, 3.4, [f"{a}: {b}" for a, b in K.WHY] + [K.WHY_HONEST], 11, False, "#000000", min_pt=7.5, bullets=True, space=5)
sl = new_slide("6. Sumber data: nyata dan asumsi")
add_table(sl, K.PROFILE_TABLE, .35, 1.05, [3.0, 6.0, 3.6], 6.0, 10)
sl = new_slide("6. Asumsi simulasi")
add_table(sl, K.ASSUME, .35, 1.05, [2.6, 6.8, 3.2], 5.6, 10.5)
add_text(sl, .35, 6.6, 12.6, .45, [K.ASSUME_NOTE], 9.5, False, GREY, min_pt=7)
sl = new_slide("6. Perhitungan yang dapat diperiksa (1/2): harga, lantai, konsesi")
add_text(sl, .45, 1.1, 6.1, 5.9, [(f"{t}: {x}", {}) for t, x in K.WORKED[:2]], 12, min_pt=7, space=8)
add_text(sl, 6.75, 1.1, 6.1, 5.9, [(f"{t}: {x}", {}) for t, x in K.WORKED[2:4]], 12, min_pt=7, space=8)
sl = new_slide("6. Perhitungan yang dapat diperiksa (2/2): injeksi, HS, kapal")
add_text(sl, .45, 1.1, 6.1, 5.9, [(f"{t}: {x}", {}) for t, x in K.WORKED[4:6]], 12, min_pt=7, space=8)
add_text(sl, 6.75, 1.1, 6.1, 5.9, [(f"{t}: {x}", {}) for t, x in K.WORKED[6:]], 12, min_pt=7, space=8)
sl = new_slide("7. Cuplikan simulasi: keluaran konsol nyata")
add_pic(sl, FIG("fig_sim_terminal.png"), .4, 1.05, 12.5, 5.5)
add_text(sl, .45, 6.55, 12.4, .5, ["Keluaran python -m kraka_mas.demo: inquiry, empat arm pada bulan yang sama, pesan antaragen pada bus."], 11, False, GREY, min_pt=8)
sl = new_slide("7. Aplikasi web: Products dan RFQ")
add_pic(sl, K.SCREENS / "01_products.png", .3, 1.05, 5.8, 5.95); add_pic(sl, K.SCREENS / "02_rfq_price_negotiation.png", 6.3, 1.05, 6.7, 5.95)
sl = new_slide("7. Aplikasi web: Procurement, Lead finder, Today")
add_pic(sl, K.SCREENS / "04_order_procurement.png", .3, 1.05, 4.2, 5.95); add_pic(sl, K.SCREENS / "06_lead_detail.png", 4.6, 1.05, 4.2, 5.95); add_pic(sl, K.SCREENS / "07_today.png", 8.9, 1.05, 4.2, 5.95)
sl = new_slide("8.1 Rancangan eksperimen: empat konfigurasi")
add_text(sl, .45, 1.1, 12.4, 5.9, ["Manual: pemilik atau admin mengerjakan semuanya (email, WhatsApp, spreadsheet).", "Agen tunggal: satu agen otonom, satu konteks, registry pemasok diperbarui berkala, tanpa gerbang manusia.",
        "B2: arsitektur multi-agen penuh, gerbang persetujuan manusia dimatikan.", "MAS: arsitektur penuh dengan persetujuan manusia (yang diimplementasikan aplikasi web).",
        "MAS vs B2 mengisolasi nilai pengawasan manusia; B2 vs agen tunggal mengisolasi nilai arsitektur.", f"{K.NS} skenario (seed 0-299), satu bulan inquiry per skenario, common random numbers, bootstrap 95%; 28 tes Python dan 16 tes aplikasi."], 14, False, "#000000", min_pt=8, bullets=True, space=8)
sl = new_slide("8.2 Hasil utama")
add_pic(sl, FIG("fig_main_results.png"), .3, 1.05, 12.7, 5.95)
sl = new_slide("8.2 Hasil utama: selisih berpasangan")
add_table(sl, K.paired(), .4, 1.1, [3.6, 3.0, 3.0, 3.0], 3.0, 12)
add_text(sl, .4, 4.3, 12.5, 2.7, [f"Terhadap manual: margin {K.ci(K.PD['mas_vs_manual_margin'], lambda x: K.n(x,0,'+',True))} USD/bulan dan tepat waktu {K.ci(K.PD['mas_vs_manual_otif'], lambda x: K.n(100*x,1,'+'))} poin.",
        f"Terhadap agen tunggal dan B2: ketepatan waktu tidak berbeda nyata; yang berbeda adalah margin, dana dialihkan, dan keamanan.", "Sebagian besar selisih margin MAS terhadap B2 berasal dari pengecualian pemilik (asumsi 40% diterima)."], 13, False, "#000000", min_pt=8, bullets=True, space=6)
sl = new_slide("8.3 Ablation: kontribusi tiap komponen")
add_pic(sl, FIG("fig_ablation.png"), .3, 1.05, 12.7, 4.3)
add_text(sl, .45, 5.4, 12.4, 1.6, K.INTERPRET[:2], 11, False, "#000000", min_pt=7.5, bullets=True)
sl = new_slide("8.3 Ablation: kondisi tertekan dan trade-off")
add_text(sl, .45, 1.1, 12.4, 5.9, K.INTERPRET[2:4], 13, False, "#000000", min_pt=8, bullets=True, space=8)
sl = new_slide("8.4 Stres: pemasok gagal dan injeksi RFQ")
add_pic(sl, FIG("fig_stress.png"), .3, 1.05, 12.7, 3.5)
add_table(sl, K.injection_table(), .4, 4.7, [2.6, 2.5, 2.5, 2.5, 2.5], 2.3, 10)
sl = new_slide("8.5 Injeksi: sapuan peluang serangan pada agen tunggal")
add_table(sl, K.single_inj_table(), .4, 1.2, [5.0, 3.5, 3.5], 2.2, 13)
add_text(sl, .4, 3.7, 12.5, 3.2, ["Laju keberhasilan serangan adalah asumsi, sehingga disapu: pada 5% kerugian kecil; pada 30% dan 60% besar.", "Bukti yang tidak bergantung asumsi: tes kode aplikasi memastikan lantai harga dan negosiasi tidak dapat diturunkan oleh teks RFQ.",
        "Tidak terukur: perilaku LLM sungguhan terhadap serangan."], 14, False, "#000000", min_pt=8, bullets=True, space=8)
sl = new_slide("8.6 Sensitivitas dan skala")
add_pic(sl, FIG("fig_sensitivity.png"), .3, 1.05, 12.7, 3.3)
add_table(sl, K.scale_table(), .4, 4.5, [2.0, 3.6, 3.6, 3.6], 1.6, 10.5)
add_table(sl, K.cost_table(), .4, 6.1, [2.2, 2.6, 2.6, 2.6, 2.6], 0.9, 8.5)
sl = new_slide("8.7 Model ML dan Marketing (eksploratif)")
add_pic(sl, FIG("fig_ml.png"), .3, 1.05, 12.7, 2.9)
add_pic(sl, FIG("fig_marketing.png"), .3, 4.0, 7.3, 2.8)
add_table(sl, K.MARKETING_TABLE, 7.7, 4.1, [2.4, 1.0, 0.9, 0.9], 2.4, 9.5)
sl = new_slide("8.8 Implikasi bisnis bagi KrakaCoal")
add_table(sl, K.BIZ, .4, 1.1, [7.2, 5.3], 5.8, 12)
sl = new_slide("9. Status komponen (1/2)")
add_table(sl, K.STATUS[:14], .3, 1.05, [4.2, 2.8, 1.9, 3.8], 5.9, 10)
sl = new_slide("9. Status komponen (2/2)")
add_table(sl, [K.STATUS[0]] + K.STATUS[14:], .3, 1.05, [4.2, 2.8, 1.9, 3.8], 5.9, 10)
sl = new_slide("10. Keterbatasan dan ancaman validitas")
add_text(sl, .45, 1.1, 12.4, 5.9, K.LIMITS, 12, False, "#000000", min_pt=8, bullets=True, space=5)
sl = new_slide("11. Kesimpulan dan rencana berikutnya")
add_text(sl, .45, 1.1, 12.4, 1.7, [f"Waktu respons {K.n(K.mv('manual','ttq_h'),0)} jam -> {K.n(K.mv('mas','ttq_h'),1)} jam; margin bersih +{K.usd(K.mv('mas','margin')-K.mv('manual','margin'))} per bulan; tepat waktu +{K.n(100*(K.mv('mas','otif')-K.mv('manual','otif')),0)} poin terhadap manual.",
        "Keunggulan MAS atas agen tunggal ada pada keamanan dan fleksibilitas pemilik, bukan ketepatan waktu. Semua hasil adalah simulasi atas data sintetis."], 14, False, "#000000", min_pt=9, bullets=True, space=6)
add_text(sl, .45, 3.0, 12.4, 4.0, K.NEXT, 12.5, False, "#000000", min_pt=8, bullets=True, space=6)
sl = new_slide("12. Tabel artefak")
add_table(sl, K.ARTEFAK, .35, 1.1, [2.8, 5.4, 4.4], 5.8, 11.5)
sl = new_slide("Lampiran A. Prompt dan spesifikasi untuk menghasilkan kode (1/2)")
add_text(sl, .45, 1.1, 6.1, 5.9, [(f"{t}: {s_}", {}) for t, s_, _ in K.APPENDIX[:3]], 10.5, min_pt=6.5, space=6)
add_text(sl, 6.75, 1.1, 6.1, 5.9, [(f"{t}: {s_}", {}) for t, s_, _ in K.APPENDIX[3:6]], 10.5, min_pt=6.5, space=6)
sl = new_slide("Lampiran A. Prompt dan spesifikasi (2/2) dan uji penerimaan")
add_text(sl, .45, 1.1, 6.1, 5.9, [(f"{K.APPENDIX[6][0]}: {K.APPENDIX[6][1]}", {})] + [(f"{t} - {u}", {}) for t, _, u in K.APPENDIX[:3]], 10.5, min_pt=6.5, space=6)
add_text(sl, 6.75, 1.1, 6.1, 5.9, [(f"{t} - {u}", {}) for t, _, u in K.APPENDIX[3:]], 10.5, min_pt=6.5, space=6)
sl = new_slide("Peta materi kuliah ke implementasi")
add_table(sl, K.LECTURE_MAP, .3, 1.05, [4.9, 5.0, 2.9], 5.9, 10.5)
sl = new_slide("Referensi")
refs = [r[2] + " doi:" + r[5] for r in sorted(K.SOTA_ROWS, key=lambda r: r[2])] + ["Prof. Dr. Azhari MT. Materi kuliah AI Agentic Technology Systems for Digital Enterprise Ecosystem, Bab 1-5, UGM.",
        "PT. Kraka Coal Indonesia. krakacoal.com (Products, FAQ); lembar harga pemilik."]
add_text(sl, .45, 1.0, 12.4, 5.0, refs, 10, False, "#000000", min_pt=6.5, bullets=True, space=3)
add_text(sl, .45, 6.0, 12.4, 1.0, [f"Kode: {K.REPO} (branch {K.BRANCH})"], 11, False, "#000000", min_pt=8)

prs.save(D / "Presentasi_Tugas1.pptx")
print("pptx ok", SN[0], "slides")

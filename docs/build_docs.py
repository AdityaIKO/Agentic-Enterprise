"""Build docs/Laporan_Tugas1.pdf and docs/Presentasi_Tugas1.pptx.  python docs/build_docs.py"""
import pathlib, sys
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from content import *

D = pathlib.Path(__file__).parent; OUT = ROOT / "outputs"; TMP = D / "_formulas"; TMP.mkdir(exist_ok=True)
NAVY, TEAL, ORANGE = "#1F2A44", "#0F8B8D", "#E07A1F"

# ---- formula images (matplotlib mathtext)
for key, tex, _ in FORMULAS:
    lines = tex.split("@@")
    fig = plt.figure(figsize=(9, 0.75 * len(lines)), facecolor="white")
    for i, ln in enumerate(lines):
        fig.text(0, 1 - (i + .5) / len(lines), ln, fontsize=15, color=NAVY, va="center")
    fig.savefig(TMP / f"{key}.png", dpi=220, bbox_inches="tight", pad_inches=0.06, facecolor="white"); plt.close(fig)

# ================================================================= PDF
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)
from PIL import Image as PILImage

fd = pathlib.Path(matplotlib.get_data_path()) / "fonts/ttf"
pdfmetrics.registerFont(TTFont("DV", fd / "DejaVuSans.ttf")); pdfmetrics.registerFont(TTFont("DVB", fd / "DejaVuSans-Bold.ttf"))
pdfmetrics.registerFontFamily("DV", normal="DV", bold="DVB", italic="DV", boldItalic="DVB")
C = lambda h: colors.HexColor(h)
body = ParagraphStyle("b", fontName="DV", fontSize=9, leading=13.2, textColor=C("#1F2937"), spaceAfter=5)
small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=10.4, textColor=C("#4B5563"))
cell = ParagraphStyle("c", parent=body, fontSize=7.6, leading=10, spaceAfter=0)
cellb = ParagraphStyle("cb", parent=cell, fontName="DVB", textColor=colors.white)
h1 = ParagraphStyle("h1", fontName="DVB", fontSize=14, leading=18, textColor=C(NAVY), spaceBefore=12, spaceAfter=6)
h2 = ParagraphStyle("h2", fontName="DVB", fontSize=10.5, leading=14, textColor=C(TEAL), spaceBefore=8, spaceAfter=3)
bul = ParagraphStyle("bl", parent=body, leftIndent=12, bulletIndent=2)
P = lambda t, s=body: Paragraph(t.replace("&", "&amp;").replace("<", "&lt;") if "<b>" not in t else t.replace("&", "&amp;"), s)


def table(rows, widths, hdr=True):
    data = [[Paragraph(str(c).replace("&", "&amp;"), cellb if (hdr and i == 0) else cell) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1 if hdr else 0)
    st = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .4, C("#D1D5DB")),
          ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
    if hdr: st += [("BACKGROUND", (0, 0), (-1, 0), C(NAVY))]
    st += [("BACKGROUND", (0, i), (-1, i), C("#F3F8F8")) for i in range(2, len(rows), 2)]
    t.setStyle(TableStyle(st)); return t


def img(path, width):
    w, h = PILImage.open(path).size; return Image(str(path), width=width, height=width * h / w)


def footer(c, d):
    c.saveState(); c.setFont("DV", 7); c.setFillColor(C("#6B7280"))
    c.drawString(2 * cm, 1.2 * cm, f"{TITLE} - {NAMA}"); c.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Hal. {d.page}"); c.restoreState()


s = []
s += [Spacer(1, 2 * cm), Paragraph(TITLE, ParagraphStyle("t", fontName="DVB", fontSize=22, leading=28, textColor=C(NAVY))),
      Spacer(1, .3 * cm), P("Laporan Progres Tugas 1", ParagraphStyle("t2", parent=body, fontSize=13, textColor=C(TEAL), leading=18)),
      P(COURSE), Spacer(1, .4 * cm), P(f"Penyusun: {NAMA} ({NIM}) - pengerjaan individu"),
      P(f"Repositori kode: {REPO}"), Spacer(1, .5 * cm), img(OUT / "fig_architecture.png", 16.5 * cm), PageBreak()]

s += [P("1. Peran Anggota Tim", h1), table([["Peran", "Tanggung jawab"]] + [list(r) for r in ROLES], [4.6 * cm, 12.4 * cm]),
      Spacer(1, 4), P(ROLE_NOTE, small)]
s += [P("2. Deskripsi Problem", h1)] + [P(t) for t in PROBLEM]
s += [P("3. Ilustrasi Data dan Perhitungan Komputasi", h1), P(DATA_NOTE, small),
      P("Contoh 3 lead (fitur mentah + hasil agen ICP-Fit):", h2), table(SAMPLE_TABLE, [2.4 * cm, 3.4 * cm, 2.2 * cm, 3 * cm, 3 * cm])]
s += [Spacer(1, 4)] + [P(t) for t in WORKED_TEXT]
s += [P("Hasil perhitungan skoring:", h2), table(CALC_TABLE, [2 * cm, 4 * cm, 3.4 * cm, 2.4 * cm, 2.4 * cm])]
s += [P("Pipeline penuh dijalankan pada seluruh data sintetis (kode: src/run_demo.py). Ringkasan keluaran agen:", h2)] + \
     [Paragraph("• " + l.replace("&", "&amp;"), bul) for l in R["log"]]
s += [P("4. Rumus Notasi Matematika", h1)]
for key, _, expl in FORMULAS:
    im = img(TMP / f"{key}.png", min(15 * cm, PILImage.open(TMP / f"{key}.png").size[0] / 220 * 2.54 * 0.75 * cm))
    s += [KeepTogether([im, Spacer(1, 2), P(expl, small), Spacer(1, 4)])]
s += [P("5. Tiga Penelitian Serupa (Lead Scoring / Kualifikasi Prospek)", h1)]
for ref, note, url in RELATED_PROBLEM:
    s += [P(ref), P(note + (f" Tautan: {url}" if url else ""), small)]
s += [P("6. Intelligent Agent System", h1), P("6.1 Mengapa multi-agent, bukan single agent?", h2), P(WHY_NOT_SINGLE)]
s += [table([["Alasan", "Penjelasan"]] + [list(x) for x in WHY_MAS], [3.8 * cm, 13.2 * cm])]
s += [P("6.2 Tiga penelitian yang menggunakan agen cerdas", h2)]
for ref, note, url in RELATED_AGENT:
    s += [P(ref), P(note + f" Tautan: {url}", small)]
s += [P(CITE_NOTE, small)]
s += [PageBreak(), P("6.3 Diagram rencana sistem", h2), img(OUT / "fig_architecture.png", 17 * cm),
      P("Alur: lead masuk -> Enrichment -> ICP-Fit -> Intent -> Scoring -> Action -> CRM/Sales. Orchestrator mengatur urutan lewat blackboard bersama; "
        "hasil won/lost dikembalikan ke Scoring Agent (feedback loop) untuk retrain.", small)]
s += [P("6.4 Komponen internal tiap agen", h2), table([AGENT_HDR] + [list(a) for a in AGENTS], [2.3 * cm, 3.2 * cm, 5.2 * cm, 3.3 * cm, 3 * cm])]
s += [P("6.5 AI, ML, atau DL?", h2)] + [P(t) for t in METHOD]
s += [P("7. Hasil Awal dan Keterbatasan", h1), table(RESULTS_ROWS, [6 * cm, 5.5 * cm, 5.5 * cm])]
s += [Spacer(1, 6)] + [Paragraph("• " + t.replace("&", "&amp;"), bul) for t in RESULT_TEXT]
s += [Spacer(1, 6), Table([[img(OUT / "fig_roc.png", 8.2 * cm), img(OUT / "fig_tiers.png", 8.2 * cm)]], colWidths=[8.5 * cm, 8.5 * cm]),
      Table([[img(OUT / "fig_weights.png", 8.4 * cm)]], colWidths=[17 * cm])]
s += [P("8. Rencana Berikutnya", h1)] + [Paragraph("• " + t, bul) for t in NEXT]
s += [P("9. Tautan Kode", h1), P(f"GitHub: {REPO} (folder src/, data/, outputs/, docs/). Menjalankan: pip install -r requirements.txt; python src/run_demo.py; python src/make_figures.py.")]
SimpleDocTemplate(str(D / "Laporan_Tugas1.pdf"), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                  bottomMargin=2 * cm, title=TITLE, author=NAMA).build(s, onFirstPage=footer, onLaterPages=footer)

# ================================================================= PPTX
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN

rgb = lambda h: RGBColor.from_string(h.lstrip("#"))
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def tb(sl, x, y, w, h, text, size=14, bold=False, color="#1F2937", align=None):
    box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color); p.space_after = Pt(6)
        if align: p.alignment = align
    return box


def slide(title, notes=""):
    sl = prs.slides.add_slide(blank)
    bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.0)); bar.fill.solid(); bar.fill.fore_color.rgb = rgb(NAVY); bar.line.fill.background()
    tb(sl, .5, .2, 12.3, .7, title, 26, True, "#FFFFFF")
    tb(sl, .5, 7.05, 12, .3, f"{TITLE}  |  {NAMA}", 9, False, "#6B7280")
    if notes: sl.notes_slide.notes_text_frame.text = notes
    return sl


def pic(sl, path, x, y, w=None, h=None):
    sl.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)


def tbl(sl, rows, x, y, w, colw, size=11, rowh=.4):
    t = sl.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rowh * len(rows))).table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size); p.font.bold = (i == 0); p.font.color.rgb = rgb("#FFFFFF" if i == 0 else "#1F2937")
            c.fill.solid(); c.fill.fore_color.rgb = rgb(NAVY if i == 0 else ("#F3F8F8" if i % 2 == 0 else "#FFFFFF"))
    return t


s0 = prs.slides.add_slide(blank)
bg = s0.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = rgb(NAVY); bg.line.fill.background()
tb(s0, .8, 2.0, 11.5, 1.6, TITLE, 38, True, "#FFFFFF"); tb(s0, .8, 3.9, 11.5, .6, "Laporan Progres Tugas 1", 22, False, "#5EEAD4")
tb(s0, .8, 4.7, 11.5, 1.4, [COURSE, f"{NAMA} ({NIM}) - pengerjaan individu", REPO], 14, False, "#CBD5E1")

sl = slide("1. Peran dalam tim"); tbl(sl, [["Peran", "Tanggung jawab"]] + [list(r) for r in ROLES], .5, 1.3, 12.3, [3.6, 8.7], 13, .62)
tb(sl, .5, 5.6, 12.3, 1.2, ROLE_NOTE, 12, False, "#4B5563")
sl = slide("2. Deskripsi problem"); tb(sl, .5, 1.3, 12.3, 5.5, PROBLEM, 15)
sl = slide("3a. Ilustrasi data (sintetis)"); tb(sl, .5, 1.2, 12.3, 1.2, DATA_NOTE, 12, False, "#4B5563")
tbl(sl, SAMPLE_TABLE, .5, 2.7, 7, [1.2, 1.9, 1.2, 1.5, 1.2], 13, .5)
tb(sl, .5, 5.4, 12.3, 1.5, R["log"][:4], 12)
sl = slide("3b. Perhitungan komputasi (contoh manual)"); tb(sl, .5, 1.2, 12.3, 3.6, WORKED_TEXT, 13)
tbl(sl, CALC_TABLE, .5, 5.0, 8.5, [1.2, 2.3, 2, 1.5, 1.5], 13, .45)
sl = slide("4. Rumus & notasi (1/2)")
y = 1.2
for key, _, expl in FORMULAS[:4]:
    pic(sl, TMP / f"{key}.png", .5, y, h=.55); tb(sl, .5, y + .55, 12.3, .6, expl, 10.5, False, "#4B5563"); y += 1.4
sl = slide("4. Rumus & notasi (2/2)")
y = 1.2
for key, _, expl in FORMULAS[4:]:
    hh = 0.85 if key == "score" else .55
    pic(sl, TMP / f"{key}.png", .5, y, h=hh); tb(sl, .5, y + hh, 12.3, .6, expl, 10.5, False, "#4B5563"); y += hh + .85
sl = slide("5. Tiga penelitian serupa (lead scoring)")
tb(sl, .5, 1.2, 12.3, 5.7, sum([[r[0], "    -> " + r[1], ""] for r in RELATED_PROBLEM], []), 12)
sl = slide("6a. Mengapa multi-agent?"); tb(sl, .5, 1.2, 12.3, 1.0, WHY_NOT_SINGLE, 13)
tbl(sl, [["Alasan", "Penjelasan"]] + [list(x) for x in WHY_MAS], .5, 2.4, 12.3, [3.0, 9.3], 12, .68)
sl = slide("6b. Tiga penelitian tentang agen cerdas")
tb(sl, .5, 1.2, 12.3, 5.2, sum([[r[0], "    -> " + r[1], ""] for r in RELATED_AGENT], []), 12); tb(sl, .5, 6.4, 12.3, .6, CITE_NOTE, 9, False, "#6B7280")
sl = slide("6c. Diagram rencana sistem"); pic(sl, OUT / "fig_architecture.png", .9, 1.15, w=11.5)
sl = slide("6d. Komponen internal agen"); tbl(sl, [AGENT_HDR] + [list(a) for a in AGENTS], .3, 1.2, 12.7, [1.7, 2.4, 4.0, 2.4, 2.2], 9, .75)
sl = slide("6e. AI, ML, atau DL?"); tb(sl, .5, 1.3, 12.3, 5.5, METHOD, 15)
sl = slide("7. Hasil awal"); tbl(sl, RESULTS_ROWS, .5, 1.2, 12.3, [5, 3.7, 3.6], 12, .42)
pic(sl, OUT / "fig_roc.png", .5, 3.9, h=3.0); pic(sl, OUT / "fig_tiers.png", 4.6, 3.9, h=3.0); pic(sl, OUT / "fig_weights.png", 8.7, 3.9, h=3.0)
sl = slide("8. Keterbatasan & rencana berikutnya"); tb(sl, .5, 1.3, 12.3, 5.5, RESULT_TEXT[2:] + NEXT + [f"Kode: {REPO}"], 15)
prs.save(D / "Presentasi_Tugas1.pptx")
print("built")
